import os
import json
import csv
import openpyxl
from typing import List, Dict


from src.process_bank import process_bank_search
from src.file_readers import read_transactions_from_excel, read_transactions_from_csv
from src.masks import get_mask_account, get_mask_card_number
from src.utils import read_transactions


# Вариант 1: Интерактивный режим
def run_interactive():
    print("Программа: Привет! Добро пожаловать в программу работы с банковскими транзакциями.")

    data = None

    while True:
        print("Выберите необходимый пункт меню:")
        print("1. Получить информацию о транзакциях из JSON-файла")
        print("2. Получить информацию о транзакциях из CSV-файла")
        print("3. Получить информацию о транзакциях из XLSX-файла")

        choice = input()

        if choice == '1':
            print("Для обработки выбран JSON-файл.")
            file_path = input("Введите путь к JSON файлу:\n")
            try:
                data = load_json(file_path)
                break
            except Exception as e:
                print(f"Ошибка загрузки файла: {e}")
                continue

        elif choice == '2':
            print("Для обработки выбран CSV-файл.")
            file_path = input("Введите путь к CSV файлу:\n")
            try:
                data = load_csv(file_path)
                break
            except Exception as e:
                print(f"Ошибка загрузки файла: {e}")
                continue

        elif choice == '3':
            print("Для обработки выбран XLSX-файл.")
            file_path = input("Введите путь к XLSX файлу:\n")
            try:
                data = load_xlsx(file_path)
                break
            except Exception as e:
                print(f"Ошибка загрузки файла: {e}")
                continue

        else:
            print("Некорректный выбор. Попробуйте снова.")

    # Работа с фильтрацией по статусу
    statuses_available = ['EXECUTED', 'CANCELED', 'PENDING']

    while True:
        print("Доступные статусы:", ', '.join(statuses_available))
        status_input = input("Введите желаемый статус:\n").strip().upper()

        if status_input not in statuses_available:
            print("Некорректный статус. Попробуйте снова.")
            continue

        filtered_data = filter_by_status(data, status_input)

        if not filtered_data:
            print("Не найдено ни одной транзакции, подходящей под ваши условия фильтрации.")

            retry_choice = input("Попробовать снова? (да/нет): ").strip().lower()
            if retry_choice != 'да':
                print("Завершение программы.")
                return
            else:
                continue

        print(f"Операции отфильтрованы по статусу \"{status_input}\"")

        # Сортировка по дате?
        sort_choice = input("Отсортировать операции по дате? Да/Нет\n").strip().lower()

        if sort_choice == 'да':
            order_choice = input("Отсортировать по возрастанию или по убыванию?\n").strip().lower()

            ascending = True if order_choice in ['по возрастанию', 'по возрастанию'] else False

            filtered_data = sort_transactions(filtered_data, ascending=ascending)

        # Фильтр по валюте (рубли?)
        rub_filter_choice = input("Выводить только рублевые транзакции? Да/Нет\n").strip().lower()

        if rub_filter_choice == 'да':
            filtered_data = filter_by_amount_currency(filtered_data, currency='руб.')

        # Фильтр по слову в описании?
        search_word_choice = input(
            "Отфильтровать список транзакций по определенному слову в описании? Да/Нет\n").strip().lower()

        if search_word_choice == 'да':
            search_word = input("Введите слово для поиска:\n").strip()
            filtered_data = process_bank_search(filtered_data, search_word)

        # Вывод итоговых операций
        if not filtered_data:
            print("Не найдено ни одной транзакции, подходящей под ваши условия фильтрации.")

            retry_choice2 = input("Попробовать снова? (да/нет): ").strip().lower()
            if retry_choice2 != 'да':
                print("Завершение программы.")
                return
            else:
                continue

        print("\nРаспечатываю итоговый список транзакций...\n")

        total_ops = len(filtered_data)

        for t in filtered_data:
            date_str = t.get('date', '')
            description = t.get('description', '')
            amount_str = t.get('amount', '')

            print(f"{date_str} {description}\nСумма: {amount_str}\n")

        print(f"Всего банковских операций в выборке: {total_ops}")


# Вариант 2: Структурированный запуск с логированием и маскированием (из второго файла)
def run_structured():
    import logging

    # Настройка логгера
    logging.basicConfig(level=logging.INFO)
    utils_logger = logging.getLogger('utils_logger')

    # Пути к файлам (можно изменить на свои пути)
    base_dir = os.path.dirname(__file__)
    json_path = os.path.join(base_dir, 'data', 'operations.json')
    csv_file_path = os.path.join(base_dir, 'data', 'transactions.csv')
    excel_file_path = os.path.join(base_dir, 'data', 'transactions_excel.xlsx')

    utils_logger.info("Запуск программы")

    # Чтение транзакций из JSON файла
    try:
        transactions = read_transactions(json_path)
        utils_logger.info(f"Загружено {len(transactions)} транзакций из JSON.")
    except Exception as e:
        utils_logger.error(f"Ошибка при чтении JSON файла: {e}")
        transactions = []

    # Обработка транзакций из JSON с маскированием и логированием
    for t in transactions:
        card_masked = get_mask_card_number(t.get('card_number', ''))
        account_masked = get_mask_account(t.get('account_number', ''))
        utils_logger.debug(f"Обработана транзакция: карта={card_masked}, счет={account_masked}")

    # Чтение из CSV файла
    try:
        print("Чтение данных из CSV файла...")
        transactions_csv = read_transactions_from_csv(csv_file_path)
        print(f"Найдено {len(transactions_csv)} транзакций из CSV.")
        for transaction in transactions_csv:
            print(transaction)
    except FileNotFoundError as e:
        print(f"Ошибка: Файл не найден. {e}")
    except Exception as e:
        print(f"Произошла непредвиденная ошибка при чтении CSV: {e}")

    print("\n" + "-" * 50 + "\n")

    # Чтение из Excel файла
    try:
        print("Чтение данных из Excel файла...")
        transactions_excel = read_transactions_from_excel(excel_file_path)
        print(f"Найдено {len(transactions_excel)} транзакций из Excel.")
        for transaction in transactions_excel:
            print(transaction)
    except FileNotFoundError as e:
        print(f"Ошибка: Файл не найден. {e}")
    except Exception as e:
        print(f"Произошла непредвиденная ошибка при чтении Excel: {e}")

    utils_logger.info("Завершение работы программы")


# Основной выбор режима запуска
def main():
    mode_choice = input(
        "Выберите режим запуска:\n"
        "1 - Интерактивный режим\n"
        "2 - Структурированный запуск с логированием\n"
        "Введите 1 или 2:\n"
    ).strip()

    if mode_choice == '1':
        run_interactive()
    elif mode_choice == '2':
        run_structured()
    else:
        print("Некорректный выбор режима. Завершение.")


# Вспомогательные функции для загрузки данных и фильтрации

def load_json(file_path: str) -> List[Dict]:
    """Загружает данные из JSON файла."""
    with open(file_path, 'r', encoding='utf-8') as f:
        return json.load(f)


def load_csv(file_path: str) -> List[Dict]:
    """Загружает данные из CSV файла."""
    with open(file_path, 'r', encoding='utf-8') as f:
        reader = csv.DictReader(f)
        return list(reader)


def load_xlsx(file_path: str) -> List[Dict]:
    """Загружает данные из XLSX файла."""
    wb = openpyxl.load_workbook(file_path)
    sheet = wb.active
    data = []
    headers = [cell.value for cell in next(sheet.iter_rows(max_row=1))]
    for row in sheet.iter_rows(min_row=2):
        record = {headers[i]: row[i].value for i in range(len(headers))}
        data.append(record)


    return data


def get_status_input() -> str:
    """Запрашивает у пользователя статус операции и валидирует его."""
    valid_statuses = ['EXECUTED', 'CANCELED', 'PENDING']
    while True:
        status = input(
            "Введите статус, по которому необходимо выполнить фильтрацию.\n"
            "Доступные для фильтровки статусы: EXECUTED, CANCELED, PENDING\n"
        ).strip().upper()
        if status in valid_statuses:
            print(f"Выбран статус: {status}")  # Отладка
            return status
        else:
            print(f"Статус операции \"{status}\" недоступен.")


def filter_by_status(data: List[Dict], status: str) -> List[Dict]:
    """Фильтрует транзакции по статусу."""
    return [transaction for transaction in data if transaction.get('state', '').upper() == status]


def sort_transactions(data: List[Dict], ascending=True) -> List[Dict]:
    """Сортирует транзакции по дате."""

    def parse_date(date_str):
        from datetime import datetime
        try:
            return datetime.strptime(date_str, '%d.%m.%Y')
        except Exception:
            from datetime import datetime
            return datetime.min  # дата некорректна

    return sorted(data, key=lambda x: parse_date(x.get('date', '')), reverse=not ascending)


def filter_by_amount_currency(data: List[Dict], currency='RUB') -> List[Dict]:
    """Фильтрует транзакции по валюте."""
    filtered = []
    for t in data:
        amount_str = t.get('amount', '')
        if amount_str.endswith(currency):
            filtered.append(t)
        elif 'USD' in amount_str or 'EUR' in amount_str:
            continue
        else:
            continue
    return filtered


if __name__ == '__main__':
    main()
